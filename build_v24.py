#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Excel de importação Businessmap a partir do v2.4, SEM o Bloco 4 (APIs — S5.x)."""
import re, json
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

LINES = open('full24.txt', encoding='utf-8').read().split('\n')
PRIO_MAP = {'Must':'High','Should':'Average','Could':'Low',"Won't":'Low','Wont':'Low'}

RE_BLOCO   = re.compile(r'^\d+\.\s*Bloco\s+(\d+)\s+—\s+(.+)$', re.IGNORECASE)
RE_TRANSV  = re.compile(r'^\d+\.\s*Transversal\s*$', re.IGNORECASE)
RE_TRANSV2 = re.compile(r'^\d+\.\s*Requisitos transversais acrescentados', re.IGNORECASE)
RE_PIECE   = re.compile(r'^\d+\.\d+\.\s*(S\d\.\d|T\d)\s+—\s+(.+)$')
RE_TEAM    = re.compile(r'^Equipa:\s*(.+?)\s+·\s+\d+\s+user stories')
RE_STORY   = re.compile(r'^(.+?)\s+·\s+US-(\d+)\s+—\s+(.+?)\s+\((US-[^)]+)\)\s*.*$')
RE_SECTION = re.compile(r'^\d+\.\s')
RE_PRIO    = re.compile(r'Prioridade:\s*([A-Za-zÃ-ÿ\']+)')
RE_SUBHEAD = re.compile(r'^T\d\s+—\s')   # sub-cabeçalho dentro de S1.4 (ex.: "T4 — Internacionalização")

def is_boundary(l):
    return bool(RE_STORY.match(l) or RE_BLOCO.match(l) or RE_TRANSV.match(l)
                or RE_TRANSV2.match(l) or RE_PIECE.match(l) or RE_SECTION.match(l)
                or RE_SUBHEAD.match(l))

stories=[]; cur_block=cur_block_name=cur_piece=cur_team=cur_deliverable=None
i=0; n=len(LINES)
while i<n:
    line=LINES[i].strip()
    m=RE_BLOCO.match(line)
    if m:
        cur_block=m.group(1); cur_block_name=f"Bloco {m.group(1)} — {m.group(2).strip()}"; i+=1; continue
    if RE_TRANSV.match(line) or RE_TRANSV2.match(line):
        cur_block='T'; cur_block_name='Transversal'; i+=1; continue
    m=RE_PIECE.match(line)
    if m:
        cur_piece=m.group(1); cur_deliverable=m.group(2).strip(); i+=1; continue
    m=RE_TEAM.match(line)
    if m:
        cur_team=m.group(1).strip(); i+=1; continue
    m=RE_STORY.match(line)
    if m:
        prefix,num,name,us_id=m.groups()
        body=[]; j=i+1
        while j<n and not is_boundary(LINES[j].strip()):
            body.append(LINES[j].rstrip()); j+=1
        stories.append(dict(us_id=us_id.strip(),name=name.strip(),block=cur_block,
            block_name=cur_block_name or '',piece=cur_piece,deliverable=cur_deliverable,
            team=cur_team,body=body)); i=j; continue
    i+=1

def build(s):
    body=s['body']; narr={}; req_idx=None
    for idx,b in enumerate(body):
        bs=b.strip()
        for k in ('Como','Quero','Para'):
            if bs.startswith(k+' ') and k not in narr: narr[k]=bs[len(k):].strip()
        if bs.startswith('Requisitos:') and 'Prioridade' in bs: req_idx=idx; break
    moscow='Must'
    if req_idx is not None:
        mp=RE_PRIO.search(body[req_idx].strip())
        if mp: moscow=mp.group(1).strip().capitalize()
    rest=[b for b in body[(req_idx+1 if req_idx is not None else 0):] if b.strip()]
    d=[]
    for k in ('Como','Quero','Para'):
        if narr.get(k): d.append(f"{k} {narr[k]}")
    if rest:
        d.append(''); d.append('──────── Critérios de Aceitação (BDD) & Regras ────────')
        for b in rest:
            bs=b.strip()
            if bs=='Regras de negócio a respeitar':
                d.append(''); d.append('▸ Regras de negócio a respeitar:')
            elif bs.startswith('Cenário'):
                d.append(''); d.append(bs)
            else: d.append(bs)
    s['_desc']='\n'.join(d).strip(); s['_prio']=PRIO_MAP.get(moscow,'Average'); s['_moscow']=moscow; s['_narr']=narr
for s in stories: build(s)

# excluir Bloco 4 (APIs — peças S5.x)
kept=[s for s in stories if not s['piece'].startswith('S5.')]
excluded=[s for s in stories if s['piece'].startswith('S5.')]

from collections import Counter
print("TOTAL parsed:",len(stories))
print("por peça:",dict(Counter(s['piece'] for s in stories)))
print("EXCLUÍDAS (Bloco 4 APIs):",len(excluded),"->",sorted(set(s['piece'] for s in excluded)))
print("MANTIDAS (sem Bloco 4):",len(kept))
# S1.4 detém, por desenho, ids US-T3/US-T4 -> não é mismatch
mism=[s['us_id'] for s in stories if s['piece']!='S1.4' and re.match(r'US-(S\d\.\d|T\d)-\d+',s['us_id']).group(1)!=s['piece']]
print("id/peça mismatches:",mism)
print("blocos:",{s['block']:s['block_name'][:45] for s in stories})
print("sem narrativa:",[s['us_id'] for s in stories if not s['_narr'].get('Como')])

# ---- XLSX ----
CARD_TYPE='Deliverable (FLSA 2)'; REQ_AREA='DNI'; PT_APP='xxxx'
HEADERS=['Title','Description','Type','Priority','Workflow name','Lane','Column','PT | Application','Requesting Area']
WIDTHS={'A':46,'B':92,'C':18,'D':10,'E':20,'F':16,'G':16,'H':18,'I':18}
hf=PatternFill('solid',fgColor='1D4ED8'); hfont=Font(bold=True,color='FFFFFF',size=11)
thin=Side(style='thin',color='D0D5DD'); bd=Border(left=thin,right=thin,top=thin,bottom=thin)
top=Alignment(vertical='top',wrap_text=True)
def write_xlsx(rows,out):
    wb=Workbook(); ws=wb.active; ws.title='Cards'; ws.append(HEADERS)
    for c in ws[1]:
        c.fill=hf; c.font=hfont; c.alignment=Alignment(vertical='center',horizontal='left'); c.border=bd
    for s in rows:
        ws.append([f"[{s['block_name']}] {s['name']}", s['_desc'], CARD_TYPE, s['_prio'],
                   'Operacional Workflow','','',PT_APP,REQ_AREA])
    for col,w in WIDTHS.items(): ws.column_dimensions[col].width=w
    for r in range(2,ws.max_row+1):
        for c in ws[r]: c.alignment=top; c.border=bd
        ws.row_dimensions[r].height=90
    ws.freeze_panes='A2'; wb.save(out); print("SAVED:",out,"cards:",ws.max_row-1)

write_xlsx(stories,'Moneyball_CMP_UserStories_v2.4.xlsx')
write_xlsx(kept,'Moneyball_CMP_UserStories_v2.4_semBloco4.xlsx')
import json; json.dump(stories,open('parsed24.json','w'),ensure_ascii=False)
